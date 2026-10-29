#!/usr/bin/env python3
"""Dump one spreadsheet sheet as JSON rows (lists of cell values) to stdout.

Used by the planner's data sources (scripts/planner-sources.mjs) so each
source reads spreadsheets the same way. Handles .xlsx, .xlsx inside a .zip
(give `archive.zip::inner/path.xlsx`) and .ods (converted with LibreOffice).

    python3 scripts/planner-sheet.py file.xlsx "Sheet name or index" [max_rows]
"""
import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile

import openpyxl


def open_book(spec):
    if "::" in spec:
        archive, inner = spec.split("::", 1)
        return openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(archive).read(inner)), read_only=True, data_only=True)
    if spec.endswith(".ods"):
        out = tempfile.mkdtemp()
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out, spec], check=True, capture_output=True)
        return openpyxl.load_workbook(os.path.join(out, os.path.basename(spec)[:-4] + ".xlsx"), read_only=True, data_only=True)
    return openpyxl.load_workbook(spec, read_only=True, data_only=True)


def main(spec, sheet, max_rows=None):
    wb = open_book(spec)
    if sheet == "?":
        json.dump(wb.sheetnames, sys.stdout)
        return
    ws = wb.worksheets[int(sheet)] if sheet.lstrip("-").isdigit() else wb[sheet]
    rows = []
    for r in ws.iter_rows(values_only=True, max_row=int(max_rows) if max_rows else None):
        rows.append([c if isinstance(c, (int, float, str)) or c is None else str(c) for c in r])
    json.dump(rows, sys.stdout)


if __name__ == "__main__":
    main(*sys.argv[1:])
